"""
Tensocraft (tensocraft.com) product scraper
--------------------------------------------
Kya karta hai:
1. Main "Products & Services" page se saari category URLs nikalta hai
   (Tensile Structure, Outdoor Shade, Modular Tensile Structure, etc.)
2. Har category page pe jaake har product block ka structured data nikalta hai:
   - Product name
   - Price (approx)
   - Spec table (Minimum Order Quantity, Fabric Material, Fabric Type,
     Application, Fabric GSM, Warranty, Frame Material, Frame Coating,
     Features, Structure Type -- jo bhi table mein ho, sab key-value pairs)
   - Description paragraph
3. Sab products ko ek list mein jama karke JSON aur Excel (.xlsx) mein save karta hai.

Chalane ka tarika:
    pip install requests beautifulsoup4 lxml openpyxl
    python scrape_tensocraft.py

Output:
    tensocraft_products.json
    tensocraft_products.xlsx
"""

import json
import re
import time
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup

BASE_URL = "https://www.tensocraft.com/"
LISTING_URL = "https://www.tensocraft.com/outdoor-shade-tensile-structure.html"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    )
}

SESSION = requests.Session()
SESSION.headers.update(HEADERS)


def get_soup(url: str) -> BeautifulSoup:
    resp = SESSION.get(url, timeout=20)
    resp.raise_for_status()
    return BeautifulSoup(resp.text, "lxml")


def get_category_urls() -> list[dict]:
    """Main listing page se saari category (sub-page) URLs nikalo."""
    soup = get_soup(LISTING_URL)
    categories = []
    seen = set()

    exclude = {
        "/", "profile.html", "enquiry.html", "sitemap.html", "photos.html",
        "outdoor-shade-tensile-structure.html",
    }

    for a in soup.select("a[href]"):
        href = a.get("href", "")
        text = a.get_text(strip=True)
        if not href.endswith(".html"):
            continue
        full_url = urljoin(BASE_URL, href)
        path = full_url.replace(BASE_URL, "")
        if path in exclude or "#" in path:
            continue
        if full_url in seen:
            continue
        if path.count("/") == 0 and text:
            seen.add(full_url)
            categories.append({"name": text, "url": full_url})

    return categories


def scrape_category(cat_url: str, cat_name: str) -> list[dict]:
    """Ek category page ke andar jitne bhi products hain unko nikalo."""
    products = []
    try:
        soup = get_soup(cat_url)
    except Exception as e:
        print(f"  [ERROR] could not load {cat_url}: {e}")
        return products

    tables = soup.find_all("table")
    for table in tables:
        heading = table.find_previous(["h1", "h2", "h3"])
        name = heading.get_text(strip=True) if heading else None

        price_match = None
        prev_text = table.find_previous(string=re.compile(r"Rs\s*[\d,]+"))
        if prev_text:
            m = re.search(r"Rs\s*[\d,]+\s*/\s*[\w\s]+", prev_text)
            if m:
                price_match = m.group(0).strip()

        specs = {}
        for row in table.find_all("tr"):
            cells = row.find_all(["td", "th"])
            if len(cells) >= 2:
                key = cells[0].get_text(strip=True)
                val = cells[1].get_text(strip=True)
                if key:
                    specs[key] = val

        desc_parts = []
        for sib in table.find_next_siblings():
            if sib.name in ("h1", "h2", "h3", "table"):
                break
            if sib.name == "p":
                t = sib.get_text(strip=True)
                if t:
                    desc_parts.append(t)
        description = " ".join(desc_parts)

        if not specs and not name:
            continue

        products.append(
            {
                "category": cat_name,
                "category_url": cat_url,
                "product_name": name,
                "approx_price": price_match,
                "specs": specs,
                "description": description,
            }
        )

    return products


def main():
    print("Fetching category list...")
    categories = get_category_urls()
    print(f"Found {len(categories)} categories.\n")

    all_products = []
    for i, cat in enumerate(categories, 1):
        print(f"[{i}/{len(categories)}] Scraping: {cat['name']} -> {cat['url']}")
        prods = scrape_category(cat["url"], cat["name"])
        if not prods:
            print(f"  -> 0 products found (check manually: {cat['url']})")
        all_products.extend(prods)
        time.sleep(1)

    print(f"\nTotal products scraped: {len(all_products)}")

    with open("tensocraft_products.json", "w", encoding="utf-8") as f:
        json.dump(all_products, f, ensure_ascii=False, indent=2)
    print("Saved -> tensocraft_products.json")

    try:
        import openpyxl
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products"

        spec_keys = []
        for p in all_products:
            for k in p["specs"].keys():
                if k not in spec_keys:
                    spec_keys.append(k)

        headers = ["category", "product_name", "approx_price", "description"] + spec_keys
        ws.append(headers)

        for p in all_products:
            row = [
                p["category"],
                p["product_name"],
                p["approx_price"],
                p["description"],
            ] + [p["specs"].get(k, "") for k in spec_keys]
            ws.append(row)

        for i, _ in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = 22

        wb.save("tensocraft_products.xlsx")
        print("Saved -> tensocraft_products.xlsx")
    except ImportError:
        print("openpyxl not installed — skipping Excel export (JSON still saved).")


if __name__ == "__main__":
    main()